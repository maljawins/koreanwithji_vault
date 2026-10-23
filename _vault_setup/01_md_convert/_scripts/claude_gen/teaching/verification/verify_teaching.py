"""Focused checks using the actual failed Preston evidence; no paid calls."""
import copy
import importlib.metadata
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
import shutil
import contextlib
import io
from types import SimpleNamespace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[6]
sys.path.insert(0, str(ROOT / '_vault_setup/01_md_convert/_scripts'))
import teaching_yt_md_pipeline as p

STATE = Path(__file__).resolve().parents[1]
OLD = STATE / 'YZjWuFVSjjM_27c7dba68cb3'
AUDIO = p.read_json(OLD / 'transcript.json')
FRAMES = p.read_json(OLD / 'frames.json')
FAILED = p.read_json(OLD / 'extract_0000.json')['result']
REVIEW = p.read_json(OLD / 'source_audit_0000.json')['result']


class Replay:
    """Only the external paid boundary is substituted; real inputs are retained."""
    def __init__(self, work, repair=True):
        self.work = work
        self.media_work = OLD
        self.repair = repair
        self.requests = []

    def request(self, name, prompt, schema, images=None):
        self.requests.append((name, prompt))
        if name.startswith('extract_'):
            result = copy.deepcopy(FAILED)
            if self.repair and '_r1' in name:
                item = result['observations'][1]
                item['audio_ids'] += ['A00003', 'A00008']
                item['content'] = item['content'].replace('구두 발음 Format, 화이트보드 표기 4 MAT', '화면 표기 4MAT')
            return result
        return copy.deepcopy(REVIEW) if '_r1' not in name or not self.repair else {'passed': True, 'issues': []}


class TeachingChecks(unittest.TestCase):
    def test_real_failed_chunk_is_repaired_from_same_sources(self):
        with tempfile.TemporaryDirectory() as directory:
            replay = Replay(Path(directory))
            args = p.parse_args([])
            evidence = p.extract_evidence(replay, AUDIO, FRAMES[:40], 46.5, args)
            self.assertIn('A00003', evidence[1]['audio_ids'])
            self.assertIn('A00008', evidence[1]['audio_ids'])
            # A representative quote can stay short while all supporting IDs are attached.
            self.assertEqual(evidence[1]['quote'], FAILED['observations'][1]['quote'])
            self.assertIn(REVIEW['issues'][0], replay.requests[2][1])

    def test_repeated_source_failure_stops_at_repair_limit(self):
        with tempfile.TemporaryDirectory() as directory:
            replay = Replay(Path(directory), repair=False)
            with self.assertRaises(p.ValidationFailure):
                p.extract_evidence(replay, AUDIO, FRAMES[:40], 46.5, p.parse_args([]))
            self.assertEqual(len(replay.requests), 4)

    def test_retry_failed_is_explicit_without_resetting_status(self):
        book = p.load_book(Path(__file__).parent / 'input_before.xlsx')
        args = p.parse_args(['--retry-failed', '--url', 'https://youtu.be/YZjWuFVSjjM'])
        jobs = p.jobs_for(book, args.url, 1, retry_failed=args.retry_failed)
        self.assertEqual(jobs[0][0].status, 'F')

    def test_unknown_ids_and_modified_quote_are_rejected(self):
        for field, bad in [('audio_ids', ['A99999']), ('frame_ids', ['V99999999']), ('quote', 'NOT IN THE SOURCE')]:
            with self.subTest(field=field):
                result = copy.deepcopy(FAILED)
                result['observations'][0][field] = bad
                issues, _ = p.check_observations(result, AUDIO[:9], FRAMES[:40], 0, 46.5, 0)
                self.assertTrue(issues)

    def test_quote_cannot_be_made_by_joining_unrelated_source_spans(self):
        result = copy.deepcopy(FAILED)
        result['observations'][0]['audio_ids'] = ['A00001', 'A00008']
        result['observations'][0]['quote'] = 'losing the room, let me show you'
        issues, _ = p.check_observations(result, AUDIO[:9], FRAMES[:40], 0, 46.5, 0)
        self.assertTrue(issues)

    def test_workbook_changed_row_is_preserved(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            path = folder / 'input.xlsx'
            shutil.copy2(Path(__file__).parent / 'input_before.xlsx', path)
            book = p.load_book(path)
            wb = openpyxl.load_workbook(path)
            wb[book.sheet].cell(2, book.columns['영상 타이틀'], 'Ji changed the title')
            wb.save(path)
            wb.close()
            before = p.digest(path)
            with self.assertRaises(p.WorkbookConflict):
                p.update_status(book, [book.entries[0]], 'O', folder)
            self.assertEqual(before, p.digest(path))

    def test_saved_result_reconciles_without_any_paid_client(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            path = folder / 'input.xlsx'
            shutil.copy2(Path(__file__).parent / 'input_before.xlsx', path)
            book = p.load_book(path)
            group = [book.entries[0]]
            args = p.parse_args(['--retry-failed', '--url', group[0].url])
            output = folder / 'saved.md'
            output.write_text('saved, validated result', encoding='utf-8')
            manifest = {group[0].vid: {'signature': p.job_signature(group[0], args, 'guide'),
                        'status': 'O', 'output': str(output), 'sha256': p.digest(output)}}
            media = folder / 'media'
            media.mkdir()
            for name, value in [('metadata.json', {}), ('transcript.json', []), ('frames.json', [])]:
                p.atomic_json(media / name, value)
            (media / 'source.mp4').write_bytes(b'fixture')
            manifest[group[0].vid].update(media_work=str(media), media_fingerprint=p.media_fingerprint(media))
            # No credentials are present and creating a paid client would fail.
            status = p.process(book, group, args, {'state': folder}, 'guide', manifest)
            self.assertEqual(status, 'O')
            after = p.load_book(path)
            self.assertEqual([e.status for e in after.entries], ['O', 'X'])

    def test_locked_workbook_is_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            path = folder / 'input.xlsx'
            shutil.copy2(Path(__file__).parent / 'input_before.xlsx', path)
            book = p.load_book(path)
            path.with_name('~$input.xlsx').touch()
            before = p.digest(path)
            with self.assertRaises(p.WorkbookConflict):
                p.update_status(book, [book.entries[0]], 'O', folder)
            self.assertEqual(before, p.digest(path))

    def test_changed_workbook_blocks_before_media_or_paid_work(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            path = folder / 'input.xlsx'
            shutil.copy2(Path(__file__).parent / 'input_before.xlsx', path)
            book = p.load_book(path)
            wb = openpyxl.load_workbook(path)
            wb[book.sheet].cell(2, book.columns['변환완료여부'], 'O')
            wb.save(path)
            wb.close()
            with self.assertRaises(p.WorkbookConflict):
                p.process(book, [book.entries[0]], p.parse_args(['--retry-failed', '--url', book.entries[0].url]), {}, 'guide', {})

    def test_response_receipt_prevents_paid_repetition_after_cache_loss(self):
        from google.genai import types
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            client = p.Gemini.__new__(p.Gemini)
            client.work = client.media_work = folder
            client.types, client.guide = types, 'guide'
            client.guard = None
            client.args = p.parse_args(['--request-interval', '0.001'])
            client.calls, client.last_call = 0, 0
            def answer(**kwargs):
                return SimpleNamespace(usage_metadata=None, text='{"ok":true}', candidates=[SimpleNamespace(finish_reason='STOP')])
            client.client = SimpleNamespace(models=SimpleNamespace(generate_content=answer))
            self.assertEqual(client.request('sample', 'prompt', p.AUDIT_SCHEMA), {'ok': True})
            (folder / 'sample.json').unlink()
            def forbidden(**kwargs):
                self.fail('completed request was sent twice')
            client.client.models.generate_content = forbidden
            self.assertEqual(client.request('sample', 'prompt', p.AUDIT_SCHEMA), {'ok': True})
            self.assertEqual(client.calls, 1)

    def test_uncertain_transport_failure_is_not_automatically_resent(self):
        from google.genai import types
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            client = p.Gemini.__new__(p.Gemini)
            client.work = client.media_work = folder
            client.types, client.guide, client.guard = types, 'guide', None
            client.args = p.parse_args(['--request-interval', '0.001'])
            client.calls, client.last_call = 0, 0
            def disconnected(**kwargs):
                raise RuntimeError('response status unknown')
            client.client = SimpleNamespace(models=SimpleNamespace(generate_content=disconnected))
            with self.assertRaises(RuntimeError):
                client.request('sample', 'prompt', p.AUDIT_SCHEMA)
            with self.assertRaises(p.PipelineError):
                client.request('sample', 'prompt', p.AUDIT_SCHEMA)
            self.assertEqual(client.calls, 1)

    def test_main_reconciles_offline_even_if_credentials_file_is_unavailable(self):
        import logging
        with tempfile.TemporaryDirectory() as directory:
            vault = Path(directory)
            base = vault / '_vault_setup/01_md_convert'
            state = base / '_scripts/claude_gen/teaching'
            state.mkdir(parents=True)
            (base / 'result_check/convert_O').mkdir(parents=True)
            (base / 'result_check/convert_F').mkdir(parents=True)
            shutil.copy2(Path(__file__).parent / 'input_before.xlsx', base / 'teaching_yt_md_convert.xlsx')
            (base / 'teaching_yt_md_guide.md').write_text('guide', encoding='utf-8')
            book = p.load_book(base / 'teaching_yt_md_convert.xlsx')
            entry = book.entries[0]
            argv = ['--execute', '--vault', str(vault), '--retry-failed', '--url', entry.url,
                    '--env-file', str(vault / 'unavailable.env')]
            args = p.parse_args(argv)
            output = base / 'result_check/convert_O/saved.md'
            output.write_text('previously saved', encoding='utf-8')
            media = state / 'media'
            media.mkdir()
            for name, value in [('metadata.json', {}), ('transcript.json', []), ('frames.json', [])]:
                p.atomic_json(media / name, value)
            (media / 'source.mp4').write_bytes(b'fixture')
            p.atomic_json(state / 'manifest.json', {entry.vid: {
                'signature': p.job_signature(entry, args, 'guide'), 'status': 'O', 'output': str(output),
                'sha256': p.digest(output), 'media_work': str(media), 'media_fingerprint': p.media_fingerprint(media)}})
            root_logger = logging.getLogger()
            original_handlers, original_level = list(root_logger.handlers), root_logger.level
            try:
                with patch.object(p, 'prepare_generation', side_effect=AssertionError('must stay offline')), contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(p.main(argv), 0)
                self.assertEqual(p.load_book(book.path).entries[0].status, 'O')
            finally:
                for handler in list(root_logger.handlers):
                    if handler not in original_handlers:
                        root_logger.removeHandler(handler)
                        handler.close()
                root_logger.setLevel(original_level)

    def test_monthly_spend_cap_is_not_retried_as_transient_429(self):
        from google.genai import types
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            client = p.Gemini.__new__(p.Gemini)
            client.work = client.media_work = folder
            client.types, client.guide, client.guard = types, 'guide', None
            client.args = p.parse_args(['--request-interval', '0.001'])
            client.calls, client.last_call = 0, 0
            class SpendCap(Exception):
                code = 429
            def capped(**kwargs):
                raise SpendCap('Your project has exceeded its monthly spending cap.')
            client.client = SimpleNamespace(models=SimpleNamespace(generate_content=capped))
            with patch.object(p.time, 'sleep'):
                with self.assertRaises(SpendCap):
                    client.request('sample', 'prompt', p.AUDIT_SCHEMA)
            self.assertEqual(client.calls, 1)

    def test_call_limit_survives_process_restart(self):
        from google import genai
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            (folder / 'calls.jsonl').write_text('{"name":"previous_attempt"}\n', encoding='utf-8')
            args = p.parse_args(['--max-calls', '1'])
            with patch.object(genai, 'Client', return_value=SimpleNamespace(close=lambda: None)):
                client = p.Gemini(args, 'guide', folder)
            with self.assertRaises(p.PipelineError):
                client.request('new', 'prompt', p.AUDIT_SCHEMA)
            self.assertEqual(client.calls, 1)


if __name__ == '__main__':
    if '--reproduce' in sys.argv:
        args = p.parse_args([])
        with tempfile.TemporaryDirectory() as directory:
            try:
                p.extract_evidence(Replay(Path(directory), repair=False), AUDIO, FRAMES[:40], 46.5, args)
            except p.ValidationFailure as exc:
                print('REPRODUCED:', exc)
        book = p.load_book(ROOT / '_vault_setup/01_md_convert/teaching_yt_md_convert.xlsx')
        print('ACTUAL ROWS:', [(e.row, e.channel, e.vid, e.status) for e in book.entries])
        print('MEDIA:', len(AUDIO), 'audio segments;', len(FRAMES), 'frames')
        print('VERSIONS:', {x: importlib.metadata.version(x) for x in ['faster-whisper', 'av', 'ctranslate2', 'google-genai', 'yt-dlp', 'python-dotenv']})
        print('OCR INSTALLED:', {x: importlib.util.find_spec(x) is not None for x in ['paddle', 'paddleocr']})
    else:
        unittest.main(verbosity=2)
