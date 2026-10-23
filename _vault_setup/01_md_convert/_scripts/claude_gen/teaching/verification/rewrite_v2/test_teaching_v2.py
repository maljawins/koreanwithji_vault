"""Offline behavior checks. No real Gemini requests or workbook edits outside tempdir."""
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

SCRIPT = Path(__file__).resolve().parents[4] / 'teaching_yt_md_pipeline.py'
spec = importlib.util.spec_from_file_location('teaching_v2', SCRIPT)
p = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = p
spec.loader.exec_module(p)

BODY = '''## 핵심 요약
설명 순서를 제시한다. 질문으로 출발한다. 실제 사례로 확인한다.

## 설명 방법
### 질문으로 시작 [00:01]
발표자는 먼저 질문하라고 주장한다.

## 영상 속 사례
### 상자 시범 [00:09]
발표자는 상자를 보여준다.

## 확인 필요
화면의 작은 글자는 읽기 어렵다.
'''
SEGMENTS = [{'start':1., 'end':4., 'text':'Ask why first.'},
            {'start':9., 'end':11., 'text':'Show the box.'}]


class FakeGemini:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.prompts = []

    def request(self, name, prompt, schema=None, images=None):
        self.prompts.append((name, prompt))
        return next(self.responses)


class TeachingTests(unittest.TestCase):
    def test_timeline_keeps_speech_before_first_and_after_last_image(self):
        self.assertTrue(hasattr(p, 'build_timeline'), 'whole-source timeline is missing')
        if hasattr(p, 'build_timeline'):
            text = p.build_timeline(SEGMENTS, [{'time':5., 'ocr':'BOX'}], 12.)
            self.assertEqual(text.count('Ask why first.'), 1)
            self.assertEqual(text.count('Show the box.'), 1)
            self.assertLess(text.index('Ask why first.'), text.index('BOX'))
            self.assertLess(text.index('BOX'), text.index('Show the box.'))

    def test_visual_selection_covers_whole_duration_with_bounded_count(self):
        self.assertTrue(hasattr(p, 'select_keyframe_times'), 'bounded visual selection is missing')
        if hasattr(p, 'select_keyframe_times'):
            probes = [{'time':i/2, 'score':255. if i%3==0 else 1.} for i in range(2000)]
            chosen = p.select_keyframe_times(probes, 60)
            self.assertLessEqual(len(chosen), 60)
            self.assertEqual(chosen[0], 0.)
            self.assertEqual(chosen[-1], 999.5)
            self.assertLessEqual(max(b-a for a,b in zip(chosen,chosen[1:])), 36.)

    def test_honest_uncertainty_passes_and_invented_source_position_fails(self):
        self.assertTrue(hasattr(p, 'validate_md'), 'Teaching Markdown validation is missing')
        if hasattr(p, 'validate_md'):
            self.assertEqual(p.validate_md(BODY, SEGMENTS, [], 12.), [])
            self.assertTrue(p.validate_md(BODY.replace('[00:09]', '[00:08]'), SEGMENTS, [], 12.))
            self.assertTrue(p.validate_md(BODY.replace('[00:09]', '[12:30]'), SEGMENTS, [], 12.))

    def test_no_demonstration_does_not_force_an_invented_example(self):
        no_example = BODY.replace('### 상자 시범 [00:09]\n발표자는 상자를 보여준다.',
                                  '영상에 실제 시범 사례가 제시되지 않았다.')
        self.assertEqual(p.validate_md(no_example,SEGMENTS,[],12.), [])

    def test_generation_uses_full_source_without_batch_quote_graph(self):
        self.assertTrue(hasattr(p, 'convert_document'), 'whole-document conversion is missing')
        if hasattr(p, 'convert_document'):
            client = FakeGemini([BODY, {'passed':True, 'issues':[]}])
            doc, issues = p.convert_document(client, SEGMENTS, [], 12., max_repairs=1)
            self.assertEqual(doc, BODY)
            self.assertEqual(issues, [])
            self.assertEqual(len(client.prompts), 2)
            for _, prompt in client.prompts:
                self.assertIn('Ask why first.', prompt)
                self.assertIn('Show the box.', prompt)

    def test_repair_is_once_and_receives_actual_audit_feedback(self):
        self.assertTrue(hasattr(p, 'convert_document'), 'bounded document repair is missing')
        if hasattr(p, 'convert_document'):
            client = FakeGemini([BODY, {'passed':False,'issues':['box example is incomplete']},
                                 BODY, {'passed':False,'issues':['still incomplete']}])
            doc, issues = p.convert_document(client, SEGMENTS, [], 12., max_repairs=1)
            self.assertIn('still incomplete', issues)
            self.assertIn('box example is incomplete', client.prompts[2][1])
            self.assertEqual(len(client.prompts), 4)
            self.assertIn('상자 시범', doc)

    def test_download_rejects_high_resolution_cache(self):
        entry = p.Entry(2,'Speaker','05_teaching','Title','https://youtu.be/abcdefghijk','X','abcdefghijk')
        with tempfile.TemporaryDirectory() as d:
            work = Path(d)
            (work/'metadata.json').write_text(json.dumps({'id':entry.vid,'title':'Title','channel':'Speaker','duration':12}),encoding='utf-8')
            (work/'source.mp4').write_bytes(b'video')
            probe = {'format':{'duration':'12'},'streams':[{'codec_type':'video','height':1080},{'codec_type':'audio'}]}
            with patch.object(p,'run_command',return_value=json.dumps(probe)):
                with self.assertRaises(p.PipelineError):
                    p.download(entry,work,120)

    def test_whole_document_receipt_survives_cache_loss_without_a_second_paid_request(self):
        from google.genai import types
        with tempfile.TemporaryDirectory() as d:
            client = p.Gemini.__new__(p.Gemini)
            client.work = client.media_work = Path(d)
            client.types, client.guide, client.guard = types, 'guide', None
            client.args = p.parse_args(['--request-interval','0.001'])
            client.calls, client.last_call = 0, 0
            def answer(**kwargs):
                self.assertEqual(kwargs['config'].response_mime_type, 'text/plain')
                return SimpleNamespace(usage_metadata=None,text=BODY,candidates=[SimpleNamespace(finish_reason='STOP')])
            client.client = SimpleNamespace(models=SimpleNamespace(generate_content=answer))
            self.assertEqual(client.request('document','prompt'), BODY.strip())
            (Path(d)/'document.json').unlink()
            def forbidden(**kwargs):
                self.fail('completed request was sent twice')
            client.client.models.generate_content = forbidden
            self.assertEqual(client.request('document','prompt'), BODY.strip())
            self.assertEqual(client.calls, 1)

    def test_uncertain_transport_is_not_resent_and_call_cap_is_persistent(self):
        from google.genai import types
        with tempfile.TemporaryDirectory() as d:
            client = p.Gemini.__new__(p.Gemini)
            client.work = client.media_work = Path(d)
            client.types, client.guide, client.guard = types, 'guide', None
            client.args = p.parse_args(['--request-interval','0.001','--max-calls','1'])
            client.calls, client.last_call = 0, 0
            def disconnected(**kwargs):
                raise RuntimeError('status unknown')
            client.client = SimpleNamespace(models=SimpleNamespace(generate_content=disconnected))
            with self.assertRaises(RuntimeError):
                client.request('document','prompt')
            with self.assertRaises(p.PipelineError):
                client.request('document','prompt')
            with self.assertRaises(p.PipelineError):
                client.request('document','different prompt')
            self.assertEqual(len((Path(d)/'calls.jsonl').read_text().splitlines()), 1)
            # Reconstruct real constructor with SDK creation substituted, no network.
            with patch('google.genai.Client',return_value=client.client):
                restarted = p.Gemini(client.args,'guide',Path(d))
            self.assertEqual(restarted.calls, 1)
            with self.assertRaises(p.PipelineError):
                restarted.request('document','third prompt')

    def make_book(self, folder, status='X'):
        import openpyxl
        path = folder/'input.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['자료구분','채널','저장위치','영상 타이틀','링크','변환완료여부'])
        ws.append(['teaching','Speaker','05_teaching','Title','https://youtu.be/abcdefghijk',status])
        wb.save(path)
        wb.close()
        return p.load_book(path)

    def test_changed_or_locked_workbook_blocks_paid_work_and_preserves_file(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            book = self.make_book(folder)
            lock = folder/'~$input.xlsx'
            lock.touch()
            before = p.digest(book.path)
            with self.assertRaises(p.WorkbookConflict):
                p.check_rows(book,book.entries)
            with self.assertRaises(p.WorkbookConflict):
                p.update_status(book,book.entries,'O',folder)
            self.assertEqual(p.digest(book.path), before)
            lock.unlink()
            wb = openpyxl.load_workbook(book.path)
            wb.active.cell(2,4,'Changed by Ji')
            wb.save(book.path)
            wb.close()
            before = p.digest(book.path)
            with self.assertRaises(p.WorkbookConflict):
                p.check_rows(book,book.entries)
            with self.assertRaises(p.WorkbookConflict):
                p.update_status(book,book.entries,'O',folder)
            self.assertEqual(p.digest(book.path), before)

    def test_failed_result_with_pending_workbook_sync_recovers_offline(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            book = self.make_book(folder,'F')
            args = p.parse_args(['--retry-failed','--url',book.entries[0].url])
            output = folder/'saved.md'
            output.write_text('existing failed draft',encoding='utf-8')
            media = folder/'media'
            media.mkdir()
            for name,value in [('metadata.json',{}),('transcript.json',[]),('frames.json',[])]:
                p.atomic_json(media/name,value)
            (media/'source.mp4').write_bytes(b'fixture')
            record = {'signature':p.job_signature(book.entries[0],args,'guide'),'status':'F',
                      'output':str(output),'sha256':p.digest(output),'media_work':str(media),
                      'media_fingerprint':p.media_fingerprint(media),'workbook_synced':False,'response_work':str(folder)}
            manifest = {book.entries[0].vid:record}
            with patch.object(p,'prepare_generation',side_effect=AssertionError('must be offline')):
                result = p.process(book,book.entries,args,{'state':folder,'manifest':folder/'manifest.json'},
                                   'guide',manifest,p.PipelineWorkers())
            self.assertEqual(result, 'F')
            self.assertTrue(manifest[book.entries[0].vid]['workbook_synced'])
            self.assertTrue(p.read_json(folder/'validation.json')['workbook_synced'])

    def test_paddle2_missing_or_too_new_setuptools_is_caught_before_execution(self):
        from importlib.metadata import PackageNotFoundError
        for installed, missing in [(None,True),('84.2.0',True),('79.0.1',False)]:
            def metadata(name):
                if name == 'paddlepaddle':
                    return '2.6.2'
                if installed is None:
                    raise PackageNotFoundError(name)
                return installed
            with self.subTest(installed=installed), patch('importlib.metadata.version',side_effect=metadata):
                issues = p.prerequisites(verbose=False)
                self.assertEqual(any('setuptools' in issue for issue in issues), missing)

    def test_recovery_write_permission_error_stops_shared_workbook_batch(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            book = self.make_book(folder)
            args = p.parse_args([])
            output = folder/'saved.md'
            output.write_text('saved',encoding='utf-8')
            media = folder/'media'
            media.mkdir()
            for name,value in [('metadata.json',{}),('transcript.json',[]),('frames.json',[])]:
                p.atomic_json(media/name,value)
            (media/'source.mp4').write_bytes(b'fixture')
            record = {'signature':p.job_signature(book.entries[0],args,'guide'),'status':'O',
                      'output':str(output),'sha256':p.digest(output),'media_work':str(media),
                      'media_fingerprint':p.media_fingerprint(media),'workbook_synced':False}
            with patch.object(p,'update_status',side_effect=PermissionError('file is held')):
                with self.assertRaises(p.WorkbookConflict):
                    p.process(book,book.entries,args,{'state':folder,'manifest':folder/'manifest.json'},
                              'guide',{book.entries[0].vid:record},p.PipelineWorkers())

    def test_full_transaction_preserves_failed_draft_and_retries_using_feedback(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            book = self.make_book(folder)
            entry = book.entries[0]
            args = p.parse_args(['--request-interval','0.001'])
            args._generation_ready = True
            paths = {'state':folder,'manifest':folder/'manifest.json','O':folder/'O','F':folder/'F'}
            paths['O'].mkdir()
            paths['F'].mkdir()
            # Local-stage fixtures only. The real paid journal, rendering, files and xlsx are exercised.
            settings = {k:getattr(args,k) for k in ('whisper_model','language','device','sample_fps','max_images')}
            key = p.sha(json.dumps({'version':2,'id':entry.vid,'settings':settings},sort_keys=True).encode())
            media = folder/f'media_v2_{entry.vid}_{key[:12]}'
            media.mkdir()
            (media/'source.mp4').write_bytes(b'local video fixture')
            (media/'frame.jpg').write_bytes(b'local image fixture')
            p.atomic_json(media/'metadata.json',{'id':entry.vid,'title':entry.title,'channel':entry.channel,'duration':12.})
            p.atomic_json(media/'transcript.json',SEGMENTS)
            p.atomic_json(media/'frames.json',[{'id':'SCREEN_000','time':0.,'file':'frame.jpg','ocr':'BOX'}])
            repaired = BODY.replace('발표자는 상자를 보여준다.','발표자는 상자를 열어 내용물을 보여준다.')
            responses = iter([BODY, {'passed':False,'issues':['box example needs its sequence']},
                              BODY, {'passed':False,'issues':['box example needs its sequence']},
                              repaired, {'passed':True,'issues':[]}])
            sent_prompts = []
            def answer(**kwargs):
                sent_prompts.append(kwargs['contents'][0].parts[0].text)
                value = next(responses)
                return SimpleNamespace(usage_metadata=None, text=json.dumps(value) if isinstance(value,dict) else value,
                                       candidates=[SimpleNamespace(finish_reason='STOP')])
            sdk = SimpleNamespace(models=SimpleNamespace(generate_content=answer),close=lambda:None)
            probe = {'format':{'duration':'12'},'streams':[{'codec_type':'video','height':720},{'codec_type':'audio'}]}
            manifest = {}
            with patch('google.genai.Client',return_value=sdk), patch.object(p,'run_command',return_value=json.dumps(probe)):
                self.assertEqual(p.process(book,book.entries,args,paths,'guide',manifest,p.PipelineWorkers()),'F')
                failed = Path(manifest[entry.vid]['output'])
                failed_text = failed.read_text(encoding='utf-8')
                self.assertIn('상자 시범', failed_text)
                self.assertIn('## 변환 검증 실패', failed_text)
                self.assertEqual(p.load_book(book.path).entries[0].status,'F')
                args.retry_failed = True
                args.url = entry.url
                book = p.load_book(book.path)
                self.assertEqual(p.process(book,book.entries,args,paths,'guide',manifest,p.PipelineWorkers()),'O')
            self.assertEqual(failed.read_text(encoding='utf-8'),failed_text)
            self.assertEqual(p.load_book(book.path).entries[0].status,'O')
            self.assertIn('box example needs its sequence',sent_prompts[-2])
            self.assertTrue(list(folder.glob('history_*.json')))


if __name__ == '__main__':
    unittest.main(verbosity=2)
