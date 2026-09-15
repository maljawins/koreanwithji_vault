"""
tech_yt_md_pipeline.py

Tech 유튜브 영상(AI·코딩·개발·vault 운영) → 구조화된 .md 변환 파이프라인 (Tech 트랙 전용).
규칙서: _vault_setup\\01_md_convert\\tech_yt_md_guide.md

이 스크립트는 kor_yt_md_pipeline.py의 사본이 아니다. 파이프라인 뼈대(동시성·백오프·비용추적·
manifest 재개)만 같은 방식을 쓰고, 한국어 트랙 전용 지점(한글 유도 Whisper 프롬프트, 한글 앵커
수집, 한글 커버리지 검증, 한국어 강의용 프롬프트, 레벨·레슨 폴더 매핑)은 하나도 들어있지 않다.

[핵심 설계]
돈 드는 일(영상 보고 듣기)은 전부 로컬(RTX 4070)에서 무료로 처리하고, Gemini 3 Flash에는
"타임라인 텍스트 + distinct 프레임 이미지"만 보내 .md만 쓰게 한다.

[Korean 트랙과 다른 점]
  - 1080p 다운로드 (화면의 코드·터미널 글자가 뭉개지면 문서를 못 쓰기 때문)
  - Whisper 유도 프롬프트가 영문 기술용어를 "영문 그대로" 받아쓰게 지시한다
  - OCR 앵커는 화면의 영문 용어·명령어·파일명 (한글 토큰이 아님)
  - 출력에 ## Transcript 섹션이 없다. 있으면 검증 실패다
  - 출력은 _vault_setup\\01_md_convert\\result_check\\ 밑 두 폴더로 평면 저장. 레벨·레슨 폴더를 찾지 않는다
    (2026-08-31부터: 성공 convert_O\\, 실패 convert_F\\. Ji 요청으로 kor 트랙도 동일 구조 적용.
     2026-09-03: result_check 폴더 자체를 01_inbox에서 이 위치로 이동)

[영상 1개 처리 8단계]
  1. 중복확인  : manifest/기존 파일에 있으면 건너뜀 (xlsx 상태만 O로 갱신)
  2. 다운로드  : 1080p mp4 1개 (yt-dlp)
  3. 분리      : ffmpeg로 오디오(16k mono) + 저해상도 probe 프레임 추출
  4. 받아쓰기  : faster-whisper large-v3 + VAD (무음 구간 환각 차단)
  5. 키프레임  : probe 프레임 변화 감지 → 선별된 시점만 1080p 원본에서 재추출 → OCR
  6. 타임라인  : AUDIO 트랙 + SCREEN(OCR) 트랙을 시간순으로 병합 + 영문 앵커 수집
  7. 합성      : Gemini 3 Flash에 (타임라인 + 프레임 이미지 + 영문 앵커) → .md
  8. 검증/저장 : 규칙·커버리지 검사 → 실패 시 1회 재생성 →
                성공: result_check\\convert_O 저장 + xlsx X→O /
                실패: result_check\\convert_F 보관 + xlsx X→F

[사용법]
    python tech_yt_md_pipeline.py --dry-run     # 작업 목록만 (API 호출 없음)
    python tech_yt_md_pipeline.py --limit 1     # 시범 1개
    python tech_yt_md_pipeline.py               # 전체(X 상태) 일괄

[설치]
    pip install google-genai yt-dlp faster-whisper paddleocr pillow python-dotenv openpyxl
    ffmpeg는 PATH에 있어야 함. CUDA GPU 권장.
    .env 에 GEMINI_API_KEY 필요 (선택: YTDLP_COOKIES_FILE).
"""

import argparse
import asyncio
import io
import json
import logging
import os
import re
import subprocess
import sys
import tempfile
import time
from datetime import datetime
from pathlib import Path

import openpyxl

# 윈도우 콘솔에서 한글 깨짐 방지
if hasattr(sys.stdout, "reconfigure") and sys.stdout.encoding != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure") and sys.stderr.encoding != "utf-8":
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


# pip으로 설치된 nvidia CUDA DLL(cublas64_12.dll 등)을 ctranslate2가 찾을 수 있도록
# site-packages/nvidia/*/bin 경로를 PATH와 DLL 검색 경로에 추가
def _register_nvidia_dlls() -> None:
    try:
        import site
        for sp in site.getsitepackages():
            nvidia_root = Path(sp) / "nvidia"
            if not nvidia_root.is_dir():
                continue
            for pkg_dir in nvidia_root.iterdir():
                bin_dir = pkg_dir / "bin"
                if bin_dir.is_dir():
                    bin_str = str(bin_dir)
                    if bin_str not in os.environ.get("PATH", ""):
                        os.environ["PATH"] = bin_str + os.pathsep + os.environ.get("PATH", "")
                    if hasattr(os, "add_dll_directory"):
                        try:
                            os.add_dll_directory(bin_str)
                        except OSError:
                            pass
    except Exception:
        pass


_register_nvidia_dlls()


# ─────────────────────────── 경로 / 상수 ────────────────────────────
# 경로를 하드코딩하지 않는다. 이 스크립트 위치에서 vault 루트를 역산한다.
# 스크립트 위치: <vault>\_vault_setup\01_md_convert\_scripts\  → 루트는 parents[3]

def _detect_base_dir() -> Path:
    env = os.environ.get("BASE_DIR", "").strip()
    if env:
        return Path(env)
    return Path(__file__).resolve().parents[3]


BASE_DIR      = _detect_base_dir()
SCRIPT_DIR    = Path(__file__).resolve().parent
XLSX_PATH     = SCRIPT_DIR.parent / "tech_yt_md_convert.xlsx"
# 2026-08-31 Ji 요청으로 성공/실패를 result_check 밑 두 폴더로 분리.
# 각각 평면 저장(하위 폴더 추가 없음)이라는 원칙은 그대로 유지.
# 2026-09-03 Ji가 result_check를 01_inbox에서 _vault_setup\01_md_convert로 이동.
RESULT_CHECK_DIR = SCRIPT_DIR.parent / "result_check"
INBOX_OUT_DIR = RESULT_CHECK_DIR / "convert_O"          # 검증 통과 .md 저장 위치
FAILED_MD_DIR = RESULT_CHECK_DIR / "convert_F"          # 검증 실패 보관. 삭제 금지

GEN_DIR     = SCRIPT_DIR / "claude_gen"                # 로그·실패기록·manifest
FAIL_LOG    = GEN_DIR / "tech_변환_실패_log.md"
PROC_LOG    = GEN_DIR / "tech_pipeline_log.txt"
MANIFEST    = GEN_DIR / "tech_manifest.json"           # 완료 URL·비용 기록 (재개용)
ACTIVE_JOBS = GEN_DIR / "tech_active_jobs.json"        # 현재 처리 중인 영상 (모니터용)

# Tech exemplary output. 2026-09-03 Ji 검수 후 확정 (가이드 섹션 13 참고).
# .env의 TECH_EXEMPLARY_MD로 다른 경로를 지정하면 그쪽이 우선한다.
_DEFAULT_EXEMPLARY_MD = (
    SCRIPT_DIR.parent.parent / "02_vault_ops"
    / "tech_jaredrhod_How I Set Up Obsidian with Claude Code_260831.md"
)
_EXEMPLARY_ENV = os.environ.get("TECH_EXEMPLARY_MD", "").strip()
EXEMPLARY_MD = Path(_EXEMPLARY_ENV) if _EXEMPLARY_ENV else _DEFAULT_EXEMPLARY_MD

# ── 모델 / 합성 설정 ──
MODEL_NAME      = os.environ.get("GEMINI_MODEL", "gemini-3-flash-preview")
GEMINI_TEMP     = 0.2
THINKING_BUDGET = int(os.environ.get("GEMINI_THINKING", "2048"))

# ── 프레임(시각 캡처) 설정 ──
INTERNAL_FPS   = 2.0    # 변화 감지용 내부 샘플링(초당 N장). 화면이 0.5초만 떠도 포착
PROBE_WIDTH    = 640    # 변화 감지 전용 저해상도 폭. 감지에는 이걸로 충분하고 디스크·시간을 아낀다
SIG_SIZE       = 64     # 변화 감지 서명 해상도 (64x64 흑백)
SIG_BLOCK      = 8      # 서명을 8x8 블록으로 쪼갬 → 화면 일부만 바뀌어도 감지
BLOCK_DIFF_TH  = 6.0    # 블록 평균 밝기차가 이 값을 넘으면 "화면이 바뀌었다"
MAX_FRAMES     = 60     # 영상당 Gemini로 보낼 최대 프레임(비용·요청크기 상한)
JPEG_QUALITY   = 3      # ffmpeg -q:v (낮을수록 고화질). OCR용 원본 프레임에 적용

# Gemini 인라인 이미지 총량 상한. API 요청 전체 한도가 20MB(base64 팽창 포함)라
# 원본 바이트 기준 10MB를 넘지 않게 축소·재인코딩한다.
GEMINI_IMAGE_BUDGET = int(float(os.environ.get("GEMINI_IMAGE_BUDGET_MB", "10")) * 1024 * 1024)

# ── OCR 설정 ──
# 가이드 섹션 7은 "ko+en"을 요구하고, 섹션 2·12는 "영문 앵커 수집"을 요구한다.
# PaddleOCR에 ko+en 조합 옵션은 없다. korean 모델이 한글과 라틴 문자를 함께 인식하므로
# 인식은 korean으로 하고, 앵커(정답지)는 영문 토큰만 골라낸다.
# 터미널 영문 인식률을 더 올리고 싶으면 .env에 TECH_OCR_LANG=en 을 넣는다.
OCR_LANG = os.environ.get("TECH_OCR_LANG", "korean").strip() or "korean"

# ── 동시성 설정 ──
PIPELINE_CONC = 4       # 동시에 진행하는 영상 수 (1080p라 디스크·대역폭 부담이 커 kor보다 낮춤)
DL_CONC       = 2       # 동시 다운로드
GPU_CONC      = 1       # Whisper는 GPU 1슬롯 직렬
GEMINI_CONC   = 1       # Gemini 동시 호출 (Tier 1 RPM/TPM 쿼터 보호)
GEMINI_MIN_INTERVAL_SEC = float(os.environ.get("GEMINI_MIN_INTERVAL_SEC", "20"))

# ── 비용 설정 (Gemini 3 Flash Standard, 백만 토큰당 USD) ──
PRICE_IN        = 0.50
PRICE_CACHED_IN = 0.05
PRICE_OUT       = 3.00
KRW_PER_USD     = float(os.environ.get("KRW_PER_USD", "1350"))
# 스크립트 내 하드 상한 없음(가이드 섹션 7). 실제 한도는 Google AI Studio 콘솔에서 관리한다.
BUDGET_KRW      = float(os.environ.get("BUDGET_KRW", "999999999"))

# ── xlsx 컬럼 (0-base) ──
# 0 자료구분 / 1 채널 / 2 목적 / 3 영상 타이틀 / 4 링크 / 5 변환완료여부
XLSX_STATUS_COL = 6     # openpyxl은 1-base → F열

# Whisper 받아쓰기 유도 프롬프트 (영문 기술용어를 영문 그대로).
# korean 트랙의 "한국어 단어는 한글로" 프롬프트와 정확히 반대다. 섞이면 Agent가 에이전트로 기록된다.
_WHISPER_PROMPT = (
    "This is a technical video about AI coding tools, software development, and knowledge management. "
    "The speaker may talk in Korean, but every technical term, product name, command, file name, "
    "option flag and config key MUST be transcribed in English/Latin script exactly as spoken, "
    "never transliterated into Hangul. "
    "Examples of correct spelling: Claude Code, CLAUDE.md, MCP, Agent, subagent, skill, hook, plugin, "
    "context window, token, prompt, commit, branch, repository, GitHub, terminal, CLI, API, npx, npm, "
    "git, JSON, YAML, markdown, .env, settings.json, /compact, /clear, /init. "
    "Write Korean speech in Hangul, but keep all of the above in English."
)


class BudgetExceeded(Exception):
    """누적 비용이 한도에 도달하면 발생 → 전체 정지."""


# ─────────────────────────────── 로깅 ──────────────────────────────────

def setup_logging() -> None:
    GEN_DIR.mkdir(parents=True, exist_ok=True)
    root = logging.getLogger()
    root.setLevel(logging.INFO)
    if root.handlers:
        return
    fh = logging.FileHandler(str(PROC_LOG), encoding="utf-8", mode="a")
    fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    root.addHandler(fh)
    ch = logging.StreamHandler()
    ch.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    root.addHandler(ch)


# ─────────────────────────── 비용 추적기 ───────────────────────────────

class CostTracker:
    """Gemini usage_metadata를 누적해 USD/원화 환산. 모니터링용 기록."""

    def __init__(self, budget_krw: float):
        self.budget_krw = budget_krw
        self.in_tokens = 0
        self.cached_tokens = 0
        self.out_tokens = 0
        self.usd = 0.0
        self._lock = asyncio.Lock()

    async def add(self, usage) -> None:
        # SDK 버전별 필드명 차이에 대비해 방어적으로 읽는다
        prompt = getattr(usage, "prompt_token_count", 0) or 0
        cached = getattr(usage, "cached_content_token_count", 0) or 0
        cand   = getattr(usage, "candidates_token_count", 0) or 0
        think  = getattr(usage, "thoughts_token_count", 0) or 0
        non_cached_in = max(prompt - cached, 0)
        cost = (
            non_cached_in * PRICE_IN
            + cached * PRICE_CACHED_IN
            + (cand + think) * PRICE_OUT
        ) / 1_000_000
        async with self._lock:
            self.in_tokens += non_cached_in
            self.cached_tokens += cached
            self.out_tokens += cand + think
            self.usd += cost

    @property
    def krw(self) -> float:
        return self.usd * KRW_PER_USD

    def over_budget(self) -> bool:
        return self.krw >= self.budget_krw

    def summary(self) -> str:
        return (
            f"누적 비용 ${self.usd:.3f} (≈₩{self.krw:,.0f}) | "
            f"입력 {self.in_tokens:,} (캐시 {self.cached_tokens:,}) · 출력 {self.out_tokens:,} 토큰"
        )


# ──────────────────────── 워커 / 모델 상태 ─────────────────────────────

class PipelineWorkers:
    """무거운 모델은 지연 로딩. 자원별 세마포어 보유."""

    def __init__(self) -> None:
        self.gemini_client = None
        self._whisper = None
        self._ocr = None
        self.whisper_name = "large-v3"
        self.gpu_sem    = asyncio.Semaphore(GPU_CONC)
        self.dl_sem     = asyncio.Semaphore(DL_CONC)
        self.gemini_sem = asyncio.Semaphore(GEMINI_CONC)
        self._ocr_lock  = asyncio.Lock()
        self._gemini_pace_lock = asyncio.Lock()
        self._gemini_last_call = 0.0

    async def throttle_gemini(self) -> None:
        """호출 시작 간 최소 간격 강제 → 429(RPM/TPM 초과) 연쇄 방지."""
        async with self._gemini_pace_lock:
            now = time.monotonic()
            wait = self._gemini_last_call + GEMINI_MIN_INTERVAL_SEC - now
            if wait > 0:
                await asyncio.sleep(wait)
            self._gemini_last_call = time.monotonic()

    def init_gemini(self, api_key: str) -> None:
        from google import genai  # 지연 import
        self.gemini_client = genai.Client(api_key=api_key)

    def get_whisper(self):
        if self._whisper is None:
            from faster_whisper import WhisperModel
            try:
                logging.info(f"faster-whisper {self.whisper_name} CUDA 로딩...")
                self._whisper = WhisperModel(self.whisper_name, device="cuda", compute_type="float16")
            except Exception as e:
                logging.warning(f"CUDA 불가({e}) → CPU 폴백")
                self._whisper = WhisperModel(self.whisper_name, device="cpu", compute_type="int8")
        return self._whisper

    def get_ocr(self):
        if self._ocr is None:
            from paddleocr import PaddleOCR
            logging.info(f"PaddleOCR(lang={OCR_LANG}, CPU) 로딩...")
            # OCR은 CPU로 (GPU는 Whisper 전용 → 8GB VRAM 경합 방지).
            # PaddleOCR 3.x는 show_log/use_gpu 인자를 제거했으므로 단계적으로 폴백한다.
            #
            # enable_mkldnn=False는 필수다. 이 머신의 PaddlePaddle CPU 추론 백엔드(oneDNN)가
            # 문서방향분류·왜곡보정 등 부가 파이프라인 단계와 맞물리면 predict() 호출에서
            # NotImplementedError(ConvertPirAttribute2RuntimeAttribute)를 던지며 죽는다.
            # 이 예외는 우리 쪽 호출부(_run_ocr)의 광범위 except가 조용히 삼켜서 빈 문자열만
            # 반환하므로, 화면이 있어도 OCR 결과가 계속 ""로 나오는데 에러 로그는 안 보인다.
            # 앵커가 항상 0개로 나오면 이 문제일 가능성이 크다.
            for kwargs in [
                {"lang": OCR_LANG, "use_gpu": False, "show_log": False, "enable_mkldnn": False},
                {"lang": OCR_LANG, "use_gpu": False, "enable_mkldnn": False},
                {"lang": OCR_LANG, "enable_mkldnn": False},
                {"lang": OCR_LANG, "use_gpu": False, "show_log": False},
                {"lang": OCR_LANG, "use_gpu": False},
                {"lang": OCR_LANG},
            ]:
                try:
                    self._ocr = PaddleOCR(**kwargs)
                    break
                except (TypeError, ValueError):
                    continue
            if self._ocr is None:
                raise RuntimeError("PaddleOCR 초기화 실패: 지원되는 인자 조합 없음")
        return self._ocr


# ──────────────────────── xlsx 파서 ────────────────────────────────────

def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    # 셀 안에 줄바꿈이 들어있는 경우가 있다(예: 목적 칸에 두 줄). 그대로 두면
    # Metadata 불릿 한 줄이 두 줄로 쪼개지므로 공백 하나로 합친다.
    return re.sub(r"\s+", " ", str(v)).strip()


def is_youtube_url(url: str) -> bool:
    if "youtube.com/playlist" in url or "youtube.com/channel" in url:
        return False
    return "youtube.com" in url or "youtu.be" in url


def parse_xlsx() -> list:
    """tech_yt_md_convert.xlsx 전체 행 파싱. status(X/O/F)와 원본 엑셀 row 번호를 포함해 반환."""
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"변환 대상 xlsx 없음: {XLSX_PATH}")
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[wb.sheetnames[0]]
    entries = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        row = list(row) + [None] * (6 - len(row))
        kind, channel, purpose, video_title, url, status = row[:6]
        url = _cell_str(url)
        if not url or not is_youtube_url(url):
            continue
        entries.append({
            "kind":        _cell_str(kind),
            "channel":     _cell_str(channel),
            "purpose":     _cell_str(purpose),
            "video_title": _cell_str(video_title),
            "url":         url,
            "status":      _cell_str(status).upper(),
            "row":         row_idx,
        })
    wb.close()
    return entries


_xlsx_save_lock = asyncio.Lock()


def _update_xlsx_status_sync(row: int, status: str) -> None:
    try:
        wb = openpyxl.load_workbook(XLSX_PATH)
        ws = wb[wb.sheetnames[0]]
        ws.cell(row=row, column=XLSX_STATUS_COL).value = status
        wb.save(XLSX_PATH)
        wb.close()
    except Exception as exc:
        # Excel에서 파일이 열려있으면 저장이 실패한다. 파이프라인은 계속 간다(manifest가 진행상황 보존).
        logging.warning(f"xlsx 상태 갱신 실패(row={row} → {status}): {exc}")


async def update_xlsx_status(row: int, status: str) -> None:
    loop = asyncio.get_event_loop()
    async with _xlsx_save_lock:
        await loop.run_in_executor(None, _update_xlsx_status_sync, row, status)


async def update_xlsx_status_many(rows: list, status: str) -> None:
    for r in rows:
        await update_xlsx_status(r, status)


# ──────────────────── 파일명 (가이드 섹션 4) ────────────────────────────

# Windows 파일명 금지문자 9개. 삭제가 아니라 공백으로 치환한다(단어가 붙으면 검색이 안 됨).
_FORBIDDEN = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
MAX_FILENAME_LEN = 200


def sanitize_component(s: str) -> str:
    s = _FORBIDDEN.sub(" ", s or "")
    s = re.sub(r"\s+", " ", s)
    return s.strip().strip(".")


def make_filename(channel: str, video_title: str, convert_date: str) -> str:
    """tech_[채널]_[영상제목]_[변환일].md — 200자 초과 시 영상제목만 잘라낸다."""
    prefix = f"tech_{sanitize_component(channel)}_"
    suffix = f"_{convert_date}.md"
    title = sanitize_component(video_title)
    room = MAX_FILENAME_LEN - len(prefix) - len(suffix)
    if room < 1:
        # 채널명만으로 이미 한도를 넘는 비정상 케이스 — 채널명 쪽을 자른다
        prefix = prefix[:MAX_FILENAME_LEN - len(suffix) - 1] + "_"
        room = MAX_FILENAME_LEN - len(prefix) - len(suffix)
    if len(title) > room:
        title = title[:room].rstrip()
    return f"{prefix}{title}{suffix}"


def find_existing_md(channel: str, video_title: str) -> Path:
    """변환일만 다른 같은 영상의 .md가 result_check\\convert_O에 이미 있는지 찾는다
    (재실행 시 중복 생성 방지). make_filename으로 이름을 만든 뒤 날짜 6자리만 정규식으로 바꿔 대조한다.
    glob을 쓰지 않는 이유: 영상 제목에 [ ] ( ) 같은 문자가 흔한데 glob은 [ ]를 문자 클래스로
    해석해 매칭이 조용히 어긋난다."""
    if not INBOX_OUT_DIR.exists():
        return None
    probe = make_filename(channel, video_title, "000000")
    stem = probe[:-len("_000000.md")]
    rx = re.compile(re.escape(stem) + r"_\d{6}\.md$", re.IGNORECASE)
    for p in sorted(INBOX_OUT_DIR.iterdir()):
        if p.is_file() and rx.match(p.name):
            return p
    return None


# ───────────────────────── yt-dlp / ffmpeg ────────────────────────────

def _ytdlp_base() -> list:
    # "yt-dlp" 콘솔스크립트 exe 래퍼가 이 환경에서 깨져 있어(rc=1, 무출력) 조용히 실패한다.
    # venv 파이썬으로 yt_dlp 모듈을 직접 실행하면 정상 동작하므로 이 경로를 쓴다.
    cmd = [sys.executable, "-m", "yt_dlp", "--js-runtimes", "node"]
    cookies_file = os.environ.get("YTDLP_COOKIES_FILE", "")
    if cookies_file:
        p = Path(cookies_file)
        if not p.is_absolute():
            p = BASE_DIR / cookies_file
        cmd += ["--cookies", str(p)]
    else:
        browser = os.environ.get("YTDLP_COOKIES_BROWSER", "")
        if browser:
            cmd += ["--cookies-from-browser", browser]
    return cmd


async def _run(args: list) -> tuple:
    proc = await asyncio.create_subprocess_exec(
        *args, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
    out, err = await proc.communicate()
    return proc.returncode, out.decode("utf-8", "replace"), err.decode("utf-8", "replace")


async def fetch_metadata(url: str) -> dict:
    rc, out, err = await _run(_ytdlp_base() + ["--dump-json", "--no-playlist", url])
    if rc != 0 or not out.strip():
        raise RuntimeError(f"yt-dlp 메타데이터 실패: {err[:300]}")
    return json.loads(out)


async def download_video(url: str, video_dir: Path, dl_sem) -> Path:
    """1080p mp4 1개만 다운로드.
    korean 트랙(480p)과 다른 유일한 이유는 화면 글자다. 터미널·에디터·설정 파일이 화면에 뜨는데
    480p에서는 글씨가 뭉개져 OCR과 Gemini가 둘 다 오인식한다. 명령어 한 글자가 틀리면 못 쓰는 문서가 된다."""
    video_dir.mkdir(parents=True, exist_ok=True)
    out_tmpl = str(video_dir / "video.%(ext)s")
    cmd = _ytdlp_base() + [
        "-f", "bestvideo[height<=1080]+bestaudio/best[height<=1080]",
        "--merge-output-format", "mp4", "--no-playlist", "-o", out_tmpl, url,
    ]
    async with dl_sem:
        rc, _, err = await _run(cmd)
    if rc != 0:
        raise RuntimeError(f"yt-dlp 영상 다운로드 실패: {err[:300]}")
    cands = [p for p in video_dir.glob("video.*") if p.is_file()]
    if not cands:
        raise FileNotFoundError("다운로드된 영상 파일 없음")
    return cands[0]


def extract_audio_sync(video_path: Path, out_dir: Path) -> Path:
    """ffmpeg로 16k mono wav 추출 (Whisper 입력용)."""
    out = out_dir / "audio.wav"
    subprocess.run(
        ["ffmpeg", "-i", str(video_path), "-vn", "-ac", "1", "-ar", "16000", "-y", str(out)],
        capture_output=True, timeout=900,
    )
    if not out.exists():
        raise RuntimeError("ffmpeg 오디오 추출 실패")
    return out


def extract_probe_frames_sync(video_path: Path, frames_dir: Path, fps: float) -> list:
    """변화 감지 전용 저해상도 프레임 추출. (path, ts초) 리스트 반환.
    1080p 원본을 초당 2장 전부 뽑으면 디스크와 시간이 크게 낭비된다. 감지에는 640px면 충분하고,
    Gemini와 OCR에 쓸 고화질 프레임은 선별이 끝난 시점에만 원본에서 다시 뽑는다."""
    frames_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["ffmpeg", "-i", str(video_path), "-vf", f"fps={fps},scale={PROBE_WIDTH}:-2",
         "-q:v", "6", "-y", str(frames_dir / "p_%05d.jpg")],
        capture_output=True, timeout=1800,
    )
    files = sorted(frames_dir.glob("p_*.jpg"))
    return [(p, i / fps) for i, p in enumerate(files)]


def extract_full_frames_sync(video_path: Path, out_dir: Path, timestamps: list) -> list:
    """선별된 시점만 원본 화질로 재추출. (path, ts) 리스트 반환."""
    out_dir.mkdir(parents=True, exist_ok=True)
    out = []
    for i, ts in enumerate(timestamps):
        dest = out_dir / f"k_{i:04d}.jpg"
        subprocess.run(
            ["ffmpeg", "-ss", f"{ts:.3f}", "-i", str(video_path), "-frames:v", "1",
             "-q:v", str(JPEG_QUALITY), "-y", str(dest)],
            capture_output=True, timeout=120,
        )
        if dest.exists() and dest.stat().st_size > 0:
            out.append((dest, ts))
    return out


# ──────────────────────── Whisper (VAD) ───────────────────────────────

def whisper_transcribe_sync(model, audio_path: str) -> list:
    """VAD 필터 ON → 무음 구간 환각 차단. (start, end, text) 리스트 반환."""
    segments, info = model.transcribe(
        audio_path,
        beam_size=5,
        task="transcribe",
        initial_prompt=_WHISPER_PROMPT,
        vad_filter=True,
        vad_parameters=dict(min_silence_duration_ms=500),
    )
    logging.info(f"Whisper 감지 언어: {info.language} (prob={info.language_probability:.2f})")
    out = []
    for seg in segments:
        t = seg.text.strip()
        if t:
            out.append((float(seg.start), float(seg.end), t))
    return out


# ──────────────────────── 키프레임 선별 ───────────────────────────────

def _signature(path: Path) -> list:
    """64x64 흑백 축소 밝기값. 블록 단위 비교의 재료."""
    from PIL import Image
    img = Image.open(path).convert("L").resize((SIG_SIZE, SIG_SIZE))
    return list(img.getdata())


def _block_change(a: list, b: list) -> float:
    """8x8 블록별 평균 밝기차 중 최댓값.
    전체 평균이 아니라 블록 최댓값을 쓰는 이유: tech 영상은 터미널에 줄 하나가 추가되거나
    코드가 한 줄 스크롤되는 식으로 화면 일부만 바뀐다. 전체 평균으로는 그 변화가 묻힌다."""
    step = SIG_SIZE // SIG_BLOCK
    worst = 0.0
    for by in range(SIG_BLOCK):
        for bx in range(SIG_BLOCK):
            total = 0
            for y in range(by * step, (by + 1) * step):
                base = y * SIG_SIZE
                for x in range(bx * step, (bx + 1) * step):
                    total += abs(a[base + x] - b[base + x])
            diff = total / (step * step)
            if diff > worst:
                worst = diff
    return worst


def select_keyframe_times(probe_frames: list, max_frames: int) -> list:
    """직전 보존 프레임 대비 화면이 의미있게 바뀐 시점만 남긴다. 시점(초) 리스트 반환."""
    kept = []          # (ts, change_score)
    last_sig = None
    for path, ts in probe_frames:
        try:
            sig = _signature(path)
        except Exception:
            continue
        if last_sig is None:
            kept.append((ts, 255.0))
            last_sig = sig
            continue
        score = _block_change(sig, last_sig)
        if score <= BLOCK_DIFF_TH:
            continue   # 화면 거의 동일 → 중복, 버림
        kept.append((ts, score))
        last_sig = sig

    # 상한 초과 시: 변화량 큰 순으로 솎되 시간 순서는 유지
    # (말하는 사람 얼굴이 계속 움직이는 영상에서 프레임이 폭증하는 것을 막는다)
    if len(kept) > max_frames:
        top = sorted(kept, key=lambda x: -x[1])[:max_frames]
        kept = sorted(top, key=lambda x: x[0])
    return [ts for ts, _ in kept]


def _run_ocr(ocr, image_path: str) -> str:
    """PaddleOCR 버전별 반환 구조 차이를 흡수해 한 줄 문자열로 반환."""
    for call in (
        lambda: ocr.ocr(image_path, cls=True),
        lambda: ocr.ocr(image_path),
        lambda: ocr.predict(image_path),
    ):
        try:
            result = call()
        except (TypeError, AttributeError):
            continue
        except Exception as exc:
            # debug가 아니라 warning으로 남긴다. 이 예외를 삼키고 ""를 반환하면 앵커가
            # 전부 0개로 나오는데, 로그 레벨이 INFO인 일반 실행에서는 원인이 하나도 안 보였다.
            logging.warning(f"OCR 실패(빈 문자열로 처리) {image_path}: {exc}")
            return ""
        return _flatten_ocr(result)
    return ""


def _flatten_ocr(result) -> str:
    if not result:
        return ""
    lines = []
    for page in result:
        if page is None:
            continue
        # PaddleOCR 3.x predict(): dict에 rec_texts 키
        if isinstance(page, dict):
            lines.extend(str(t) for t in page.get("rec_texts", []))
            continue
        for item in page:
            if isinstance(item, (list, tuple)) and len(item) >= 2:
                tc = item[1]
                lines.append(str(tc[0]) if isinstance(tc, (list, tuple)) else str(tc))
            elif hasattr(item, "text"):
                lines.append(item.text)
    return " ".join(x.strip() for x in lines if x and x.strip()).strip()


def ocr_keyframes_sync(full_frames: list, ocr) -> list:
    """보존 프레임에 OCR. [(path, ts, ocr_text), ...] 반환."""
    out = []
    for path, ts in full_frames:
        out.append((path, ts, _run_ocr(ocr, str(path))))
    return out


# ──────────────────────── 영문 앵커 수집 ──────────────────────────────
# 화면에서 읽은 영문 용어·명령어·파일명 = Whisper 한글 음차를 교정할 정답지(가이드 섹션 6).
# 한글 토큰은 모으지 않는다.

_TOKEN = re.compile(r"[A-Za-z0-9_.\-/\\@]*[A-Za-z][A-Za-z0-9_.\-/\\@]*")

# 기술 표식이 없는 평범한 영어 단어는 앵커로서 가치가 없고 커버리지 검사만 어지럽힌다.
_STOPWORDS = {
    "the", "and", "for", "you", "this", "that", "with", "from", "your", "are", "can", "will",
    "not", "but", "all", "how", "what", "when", "why", "use", "using", "used", "get", "got",
    "more", "new", "one", "two", "add", "see", "run", "out", "now", "has", "have", "here",
    "there", "then", "than", "into", "just", "like", "make", "made", "want", "need", "very",
    "some", "any", "our", "its", "was", "were", "been", "they", "them", "his", "her", "who",
    "which", "about", "after", "before", "also", "only", "over", "under", "each", "such",
    "way", "time", "part", "next", "back", "good", "best", "much", "many", "even", "still",
    "let", "lets", "does", "did", "doing", "must", "should", "would", "could", "may", "might",
    "yes", "no", "ok", "okay", "click", "video", "watch", "like", "subscribe", "channel",
}

_TECH_MARK = re.compile(r"[./\\_\-]|\d")


def _is_technical(tok: str) -> bool:
    """기술 표식: 경로·확장자·스네이크·하이픈·숫자 포함, 전부 대문자, 또는 camelCase."""
    if _TECH_MARK.search(tok):
        return True
    if tok.isupper() and len(tok) >= 2:
        return True
    if re.search(r"[a-z][A-Z]", tok):
        return True
    return False


def collect_english_anchors(keyframes: list) -> tuple:
    """화면 OCR에서 영문 토큰 수집.
    반환: (프롬프트에 넣을 전체 앵커 목록, 커버리지 검사에 쓸 고신뢰 앵커 목록)"""
    freq = {}
    for _, _, text in keyframes:
        seen_here = set()
        for raw in _TOKEN.findall(text or ""):
            # 뒤쪽 문장부호는 떼되 앞의 점은 남긴다. .env / .gitignore 같은 dotfile 이름이
            # 앵커에서 사라지면 정답지 역할을 못 한다.
            tok = raw.rstrip("._-/\\@").lstrip("_-/\\@")
            if len(tok) < 2 or not re.search(r"[A-Za-z]", tok):
                continue
            low = tok.lower()
            if low in _STOPWORDS and not _is_technical(tok):
                continue
            if low in seen_here:
                continue
            seen_here.add(low)
            freq[tok] = freq.get(tok, 0) + 1

    anchors = sorted(freq, key=lambda t: (-freq[t], t.lower()))[:400]

    # 고신뢰 = 기술 표식이 있고, (경로·확장자를 포함하거나 여러 프레임에 반복 등장한 것).
    # 한 프레임에만 스친 짧은 토큰은 OCR 오인식일 확률이 높아 검증 기준에서 뺀다.
    high = []
    for tok in anchors:
        if not _is_technical(tok) or len(tok) < 3:
            continue
        if ("." in tok or "/" in tok or "\\" in tok) or freq[tok] >= 2:
            high.append(tok)
    return anchors, high[:120]


# ──────────────────────── 타임라인 병합 ───────────────────────────────

def _ts_label(sec: float, long_form: bool) -> str:
    s = int(sec)
    if long_form:
        return f"{s // 3600}:{(s % 3600) // 60:02d}:{s % 60:02d}"
    return f"{s // 60:02d}:{s % 60:02d}"


def build_timeline(segments: list, keyframes: list, long_form: bool) -> str:
    """AUDIO 트랙(Whisper)과 SCREEN 트랙(키프레임 OCR)을 시간순으로 병합.
    화면 상태 i는 [ts_i, ts_{i+1}) 구간 동안 떠 있다고 보고 그 구간의 오디오를 묶는다.
    무음 구간은 (무음)으로 명시한다 — tech 영상에서는 말없이 화면만 보여주는 구간이 곧 정보다."""
    if not keyframes:
        return "\n".join(
            f"[{_ts_label(s, long_form)}–{_ts_label(e, long_form)}] AUDIO: \"{t}\""
            for s, e, t in segments
        ) or "(내용 없음)"

    bounds = [ts for _, ts, _ in keyframes] + [float("inf")]
    lines = []
    prev_screen = None
    for i, (_, ts, ocr_text) in enumerate(keyframes):
        win_start, win_end = bounds[i], bounds[i + 1]
        spoken = [t for (s, e, t) in segments if win_start <= (s + e) / 2 < win_end]
        audio_str = " ".join(spoken) if spoken else "(무음)"
        end_label = _ts_label(win_end, long_form) if win_end != float("inf") else "끝"
        lines.append(f"[{_ts_label(win_start, long_form)}–{end_label}] AUDIO: \"{audio_str}\"")
        if ocr_text and ocr_text != prev_screen:
            lines.append(f"  SCREEN: \"{ocr_text}\"")
            prev_screen = ocr_text
        elif not ocr_text:
            lines.append("  SCREEN: (텍스트 없는 화면 — 첨부 프레임 이미지 참조)")
    return "\n".join(lines)


# ──────────────────────── 시스템 프롬프트 ─────────────────────────────

def build_system_prompt() -> str:
    """가이드 섹션 5·6·8 규칙 + (있으면) tech exemplary. 호출마다 동일 → 암묵 캐싱으로 입력비 절감."""
    exemplary_block = ""
    if EXEMPLARY_MD and EXEMPLARY_MD.exists():
        exemplary_block = (
            "\n═══ GOLD-STANDARD EXEMPLARY OUTPUT "
            "(match this depth, prose style, and code-block handling) ═══\n"
            + EXEMPLARY_MD.read_text(encoding="utf-8")
        )
    else:
        logging.warning(
            "Tech exemplary 미지정 — 규칙만으로 생성합니다. "
            "시범 변환 1개를 검수해 확정한 뒤 .env의 TECH_EXEMPLARY_MD에 경로를 넣으세요."
        )

    return f"""You convert a technical YouTube video (AI tooling, coding, software development,
knowledge-vault operation) into a structured Markdown knowledge document written in KOREAN.

The consumer of this file is an LLM agent that will later search it and apply it to real work.
That agent must be able to reproduce the described method from THIS FILE ALONE, without the video.

The quality bar is signal-to-noise, not volume. This is COMPRESSION, not preservation.
DISCARD: greetings, small talk, channel intros, subscribe/like requests, sponsorships,
self-promotion, and any point the speaker restates without adding information.
KEEP: conclusions, procedures, exact settings, exact commands, and the conditions under which a
method does or does not apply.

═══ INPUT YOU RECEIVE ═══
- [VIDEO METADATA]
- [TIMELINE]: AUDIO track (Whisper) and SCREEN track (on-screen OCR) merged in time order.
  When AUDIO is "(무음)" there is no speech; the SCREEN content IS the information for that moment.
- [KEY FRAMES]: the distinct on-screen states as images, in time order. LOOK at them directly.
  Terminal output, editor panes, config files and folder trees are only reliably readable in these
  images — the OCR text loses indentation and line breaks, which makes code unusable.
- [OCR ANCHORS]: English terms, commands, file names and config keys detected on screen.
  These are GROUND TRUTH spellings. Use them to repair the audio track.

═══ OUTPUT STRUCTURE (follow exactly) ═══
# [채널명] | [영상 제목]

## Metadata
(bullet list ONLY. A table like "| Field | Content |" is STRICTLY FORBIDDEN.)
- **채널:** ...
- **URL:** ...
- **Duration:** MM:SS
- **Upload Date:** YYYY-MM-DD
- **목적:** ...          ← copy the 목적 value from [VIDEO METADATA] verbatim
- **변환일:** YYYY-MM-DD  ← copy the 변환일 value from [VIDEO METADATA] verbatim

---

## 핵심 요약

---

## 적용 지식

### [설명적인 섹션 제목] [MM:SS]

---

## 주의사항

There is NO `## Transcript` section. Do not produce one under any name.

═══ 핵심 요약 RULES ═══
- 3 to 5 lines. Write what the video CONCLUDES, not what it is about.
  ✓ "CLAUDE.md는 짧을수록 좋다. 200줄을 넘으면 모델이 중간 내용을 무시하기 시작하므로,
     상세 규칙은 별도 파일로 빼고 CLAUDE.md에는 참조 링크만 남긴다."
  ✗ "CLAUDE.md 작성 방법을 설명하는 영상이다."  ← contains no information

═══ 적용 지식 RULES (the core of the document) ═══
- Split by topic. Section titles are DESCRIPTIVE and carry NO number prefix.
  ✓ `### CLAUDE.md를 짧게 유지해야 하는 이유 [03:40]`
  ✗ `### 1. CLAUDE.md 작성법`
- Every section title ends with a timestamp `[MM:SS]` marking where that content starts in the
  video, so a human can go back and check the source. Use `[H:MM:SS]` for videos over one hour.
  Take the timestamps from the [TIMELINE] you are given. Never invent one.
- Body is KOREAN PROSE — continuous explanatory sentences, not bullet dumps.
  Exception: information that is inherently a list (command lists, setting entries, comparison
  tables) may stay as bullets or a table.
- REAL ARTIFACTS GO IN CODE BLOCKS, VERBATIM. Commands, config file contents, folder trees, and
  prompt examples must be reproduced exactly as shown on screen inside fenced code blocks with a
  language tag. Never paraphrase them into prose.
  ```bash
  claude mcp add filesystem -- npx -y @modelcontextprotocol/server-filesystem ~/vault
  ```
- DO NOT over-summarize. Every procedure, precondition, exception and failure case the video
  presents must appear. Merge only genuine repetition of the same point.
- The bar: an agent reading this file must be able to actually perform the method. If a step's
  exact value, path, or flag was shown on screen, it belongs here.

═══ 주의사항 RULES ═══
- Collect the preconditions and limits the video EXPLICITLY states: required versions, plans,
  platforms ("Claude Code 2.0 이상", "macOS 기준 경로"), cases where the method breaks, and
  failures or trial-and-error the speaker reports first-hand.
- Never infer or invent a limitation. If the video states none, OMIT THE ENTIRE SECTION.
  Never leave an empty section behind.

═══ LANGUAGE RULES (critical) ═══
- Explanation and narration are in KOREAN prose.
- On a technical term's first appearance, give Korean plus the English original:
  `컨텍스트 창(context window)`, `토큰(token)`.
- COMMANDS, CODE, FILE NAMES, FOLDER NAMES, CONFIG KEYS, OPTION FLAGS AND PRODUCT NAMES STAY IN
  ENGLISH, EXACTLY AS WRITTEN. Never translate or transliterate them.
  ✓ `CLAUDE.md`, `MCP`, `/compact`, `npx`, `Claude Code`, `Agent`, `subagent`, `GitHub`, `commit`
  ✗ `클로드 점 엠디`, `엠씨피`, `컴팩트 명령어`, `클로드 코드`, `에이전트`, `서브에이전트`, `깃허브`
- The Whisper audio track WILL contain Hangul transliterations of English terms
  (에이전트 → Agent, 엠씨피 → MCP, 클로드 코드 → Claude Code, 깃허브 → GitHub,
   커밋 → commit, 컨텍스트 윈도우 → context window, 서브에이전트 → subagent).
  You MUST convert every one of them back to the English original. Use [OCR ANCHORS] as the
  authority for spelling. Where the anchors are silent but the English term is obvious from
  context, still write English. Only when you genuinely cannot tell, write `한글음차(English)`
  keeping both.
  This matters because the agent that later searches this vault will search for "MCP" and "Agent".
  A document written in Hangul transliteration will never be found, and is the same as not existing.

═══ VISUAL CONTENT ═══
In a tech video the SCREEN IS THE INFORMATION. When the speaker says "이렇게 설정하면 됩니다",
the actual setting exists only on screen — the audio alone cannot reconstruct it.
Read the [KEY FRAMES] images and fold everything they show into the 적용 지식 prose and code blocks.
Do NOT list frames or timestamps separately as an inventory; the code and settings you read off the
screen belong inside the relevant section's code block.

═══ DO NOT WRITE ═══
- Any "이 vault에 적용 가능한가" judgement section. You do not know this vault's structure.
  Guessing here creates statements that get re-quoted later as fact. Application decisions are made
  by the agent reading this file, not baked in at conversion time.
- Opinions about the video ("매우 유익한 영상이다").
- Subscribe/like prompts, ads, sponsorship disclosures, greetings.

═══ CHECKLIST (verify before output) ═══
[ ] Metadata is a bullet list, not a table, and includes 목적 and 변환일
[ ] `## 핵심 요약` states a conclusion, not a topic description
[ ] `## 적용 지식` section titles have no number prefix and end with a [MM:SS] timestamp
[ ] Body is Korean prose
[ ] Commands, configs and code appear verbatim in fenced code blocks
[ ] No Hangul transliteration of English technical terms remains
[ ] What was on screen is reflected in the code blocks
[ ] NO `## Transcript` section exists
[ ] No small talk, subscribe prompts, ads
[ ] No guesses and no application judgements that the video did not state
{exemplary_block}
Output ONLY the raw .md content. No code fences around the whole document, no preamble, no commentary.
"""


# ──────────────────────── Gemini 이미지 준비 ──────────────────────────

def prepare_gemini_images(keyframes: list, budget: int) -> list:
    """프레임을 인라인 전송용 JPEG 바이트로 변환.
    1080p 원본 60장은 API 요청 한도(20MB, base64 팽창 포함)를 넘길 수 있으므로
    총량이 예산 안에 들어올 때까지 해상도·품질을 단계적으로 낮춘다. 글자 판독이 목적이라
    가장 큰 단계부터 시도하고, 그래도 안 되면 마지막에 장수를 줄인다."""
    from PIL import Image

    def encode(paths, max_w, quality):
        out = []
        for p in paths:
            try:
                img = Image.open(p).convert("RGB")
            except Exception:
                continue
            if img.width > max_w:
                img = img.resize((max_w, max(1, round(img.height * max_w / img.width))),
                                 Image.LANCZOS)
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=quality, optimize=True)
            out.append(buf.getvalue())
        return out

    paths = [p for p, _, _ in keyframes]
    for max_w, quality in ((1600, 85), (1280, 80), (1024, 72)):
        blobs = encode(paths, max_w, quality)
        if sum(len(b) for b in blobs) <= budget:
            return blobs

    # 마지막 수단: 가장 작은 설정에서 예산에 들어갈 때까지 뒤에서부터 프레임을 뺀다
    blobs = encode(paths, 1024, 70)
    while blobs and sum(len(b) for b in blobs) > budget:
        blobs.pop()
    logging.warning(f"이미지 예산 초과 → 프레임 {len(keyframes)}장 중 {len(blobs)}장만 전송")
    return blobs


# ──────────────────────── Gemini 호출 ──────────────────────────────────

def _format_duration(seconds) -> str:
    s = int(seconds or 0)
    if s >= 3600:
        return f"{s // 3600}:{(s % 3600) // 60:02d}:{s % 60:02d}"
    return f"{s // 60}:{s % 60:02d}"


def call_gemini_sync(client, system_prompt, entry, metadata, timeline,
                     image_blobs, anchors, convert_date_iso, retry_note=""):
    """동기 Gemini 호출 (executor에서 실행). (text, usage) 반환."""
    from google.genai import types  # 지연 import

    duration_str = metadata.get("duration_string") or _format_duration(metadata.get("duration"))
    upload_raw   = metadata.get("upload_date", "")
    upload_date  = f"{upload_raw[:4]}-{upload_raw[4:6]}-{upload_raw[6:]}" if len(upload_raw) == 8 else upload_raw
    video_title  = entry.get("video_title") or metadata.get("title", "")
    anchors_str  = ", ".join(anchors) if anchors else "(없음)"

    retry_block = ""
    if retry_note:
        retry_block = (
            f"\n\n[RETRY NOTE — CRITICAL]\n{retry_note}\n\n"
            "The above MUST be fixed in this output. Do not drop, compress or paraphrase the listed "
            "on-screen items — reproduce them verbatim in the relevant code block or prose."
        )

    intro = f"""[VIDEO METADATA]
Channel: {entry['channel']}
URL: {entry['url']}
Duration: {duration_str}
Upload Date: {upload_date}
목적: {entry['purpose']}
변환일: {convert_date_iso}
Title: {video_title}

[TIMELINE]
{timeline}
"""

    tail = f"""[OCR ANCHORS]
{anchors_str}

[INSTRUCTION]
Generate the .md file per the system prompt. Output the .md content only, no preamble.{retry_block}
"""

    # contents 구성: intro 텍스트 → 프레임 이미지(시간순) → OCR 앵커 + 지시
    contents = [types.Part.from_text(text=intro)]
    for blob in image_blobs:
        contents.append(types.Part.from_bytes(data=blob, mime_type="image/jpeg"))
    contents.append(types.Part.from_text(text=tail))

    cfg_kwargs = dict(system_instruction=system_prompt, temperature=GEMINI_TEMP)
    if THINKING_BUDGET > 0:
        try:
            cfg_kwargs["thinking_config"] = types.ThinkingConfig(thinking_budget=THINKING_BUDGET)
        except Exception:
            pass
    config = types.GenerateContentConfig(**cfg_kwargs)

    for attempt in range(5):
        try:
            resp = client.models.generate_content(model=MODEL_NAME, contents=contents, config=config)
            return resp.text, getattr(resp, "usage_metadata", None)
        except Exception as e:
            msg = str(e)
            if "429" in msg or "RESOURCE_EXHAUSTED" in msg:
                wait = 30 * (2 ** attempt)   # 30, 60, 120, 240, 480초
                logging.warning(f"Gemini 429 → {wait}s 후 재시도 ({attempt+1}/5)")
                time.sleep(wait)
            elif "503" in msg or "UNAVAILABLE" in msg:
                wait = 15 * (2 ** attempt)   # 15, 30, 60, 120, 240초
                logging.warning(f"Gemini 503 → {wait}s 후 재시도 ({attempt+1}/5)")
                time.sleep(wait)
            else:
                raise
    raise RuntimeError("Gemini 503/429: 5회 재시도 소진")


# ──────────────────────── 검증 (가이드 섹션 11) ────────────────────────

# Whisper가 영문 기술용어를 한글로 받아쓴 흔적. 남아 있으면 나중에 검색으로 이 문서를 못 찾는다.
_TRANSLIT = [
    "에이전트", "서브에이전트", "클로드 코드", "클로드코드",
    "엠씨피", "엠시피", "깃허브", "깃헙", "깃 허브",
    "컨텍스트 윈도우", "컨텍스트윈도우", "커맨드라인", "리포지토리", "레파지토리",
]


def find_translit(content: str) -> list:
    """한글 음차 잔존 검사. `한글음차(English)` 병기 형태는 가이드가 허용하므로 통과시킨다."""
    hits = []
    for word in _TRANSLIT:
        for m in re.finditer(re.escape(word), content):
            after = content[m.end():m.end() + 3]
            if after.lstrip().startswith("("):   # 병기 형태 → 허용
                continue
            hits.append(word)
            break
    return hits


def validate_md(content: str) -> list:
    errors = []
    if "## Metadata" not in content:
        errors.append("## Metadata 섹션 누락")
    if "## 핵심 요약" not in content:
        errors.append("## 핵심 요약 섹션 누락")
    if "## 적용 지식" not in content:
        errors.append("## 적용 지식 섹션 누락")

    # tech 트랙은 Transcript를 만들지 않는다. 있으면 규칙 위반이다.
    if re.search(r"^##+\s*Transcript", content, re.MULTILINE):
        errors.append("## Transcript 섹션이 존재함 (tech 트랙 금지)")

    # Metadata가 테이블인지 (불릿 리스트만 허용)
    meta = content.split("## Metadata", 1)[-1].split("\n## ", 1)[0] if "## Metadata" in content else ""
    if re.search(r"^\s*\|", meta, re.MULTILINE):
        errors.append("Metadata가 테이블 형식임 (불릿 리스트만 허용)")
    if "**목적:**" not in meta:
        errors.append("Metadata에 목적 누락")
    if "**변환일:**" not in meta:
        errors.append("Metadata에 변환일 누락")

    body = content.split("## 적용 지식", 1)[-1] if "## 적용 지식" in content else ""
    heads = re.findall(r"^###\s+(.+)$", body, re.MULTILINE)
    if not heads:
        errors.append("적용 지식에 ### 섹션이 없음")
    else:
        if any(re.match(r"^\d+[.)]", h.strip()) for h in heads):
            errors.append("적용 지식 섹션명에 번호 prefix가 있음")
        no_ts = [h for h in heads if not re.search(r"\[\d{1,2}:\d{2}(:\d{2})?\]\s*$", h.strip())]
        if no_ts:
            errors.append(f"섹션명에 타임스탬프 없음: {no_ts[:3]}")

    # 본문이 산문인지 간이 점검 (불릿 떡칠 방지). 코드블록은 세지 않는다.
    prose_src = re.sub(r"```.*?```", "", body, flags=re.DOTALL)
    bullets = len(re.findall(r"^\s*[-*]\s+", prose_src, re.MULTILINE))
    paras = len([l for l in prose_src.splitlines()
                 if l.strip() and not l.strip().startswith(("#", "-", "*", ">", "|"))])
    if paras > 0 and bullets > paras * 1.5:
        errors.append("적용 지식이 산문보다 불릿 위주 (산문 규칙 위반 의심)")

    hits = find_translit(content)
    if hits:
        errors.append(f"한글 음차 잔존: {hits[:5]}")
    return errors


def coverage_missing(content: str, anchors: list) -> list:
    """화면에 떴던 영문 명령어·파일명 중 출력 .md에 안 들어간 것 = 시각정보 누락 신호."""
    low = content.lower()
    return [a for a in anchors if a.lower() not in low]


COVERAGE_RETRY_RATIO = 0.5   # 고신뢰 앵커의 절반 넘게 빠지면 1회 재생성


# ──────────────────────── 실패 / manifest / 진행상황 ───────────────────

def log_failure(url: str, filename: str, reason: str) -> None:
    GEN_DIR.mkdir(parents=True, exist_ok=True)
    with open(FAIL_LOG, "a", encoding="utf-8") as f:
        f.write(f"- **{filename}** (`{url}`): {reason}\n")
    logging.error(f"실패 기록: {filename} — {reason}")


_active: dict = {}


def _save_active() -> None:
    GEN_DIR.mkdir(parents=True, exist_ok=True)
    ACTIVE_JOBS.write_text(json.dumps(_active, ensure_ascii=False, indent=2), encoding="utf-8")


def _job_set(label: str, entry: dict, step: int, stage: str) -> None:
    _active[label] = {
        "label": label,
        "track": "tech",
        "channel": entry.get("channel", ""),
        "purpose": entry.get("purpose", ""),
        "step": step,
        "stage": stage,
        "ts": datetime.now().isoformat(timespec="seconds"),
    }
    _save_active()


def _job_clear(label: str) -> None:
    _active.pop(label, None)
    _save_active()


def load_manifest() -> dict:
    if MANIFEST.exists():
        try:
            return json.loads(MANIFEST.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_manifest(m: dict) -> None:
    GEN_DIR.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps(m, ensure_ascii=False, indent=2), encoding="utf-8")


# ──────────────────────── 영상 1개 파이프라인 ─────────────────────────

async def process_video(entry, filename, convert_date_iso, workers, system_prompt, tracker, manifest):
    url = entry["url"]
    target_path = INBOX_OUT_DIR / filename
    label = filename[:60]

    if tracker.over_budget():
        raise BudgetExceeded(tracker.summary())

    with tempfile.TemporaryDirectory(prefix="tech2md_") as _tmp:
        tmp = Path(_tmp)
        loop = asyncio.get_event_loop()
        try:
            _job_set(label, entry, 1, "메타데이터 수집")
            logging.info(f"[{label}] 1.메타데이터")
            metadata = await fetch_metadata(url)
            duration = float(metadata.get("duration") or 0)
            long_form = duration >= 3600

            _job_set(label, entry, 2, "영상 다운로드(1080p)")
            logging.info(f"[{label}] 2.다운로드(1080p)")
            video_path = await download_video(url, tmp / "v", workers.dl_sem)

            _job_set(label, entry, 3, "오디오/프레임 분리")
            logging.info(f"[{label}] 3.오디오/probe 프레임 분리")
            audio_path = await loop.run_in_executor(
                None, lambda: extract_audio_sync(video_path, tmp))
            probe = await loop.run_in_executor(
                None, lambda: extract_probe_frames_sync(video_path, tmp / "probe", INTERNAL_FPS))
            logging.info(f"[{label}] probe 프레임 {len(probe)}장")

            _job_set(label, entry, 4, "Whisper 받아쓰기")
            logging.info(f"[{label}] 4.Whisper(VAD)")
            async with workers.gpu_sem:   # GPU 1슬롯 직렬
                segments = await loop.run_in_executor(
                    None, lambda: whisper_transcribe_sync(workers.get_whisper(), str(audio_path)))
            if not segments and not probe:
                raise RuntimeError("오디오·화면 모두 비어있음")

            _job_set(label, entry, 5, "키프레임 선별+OCR")
            logging.info(f"[{label}] 5.키프레임 선별")
            times = await loop.run_in_executor(
                None, lambda: select_keyframe_times(probe, MAX_FRAMES))
            full = await loop.run_in_executor(
                None, lambda: extract_full_frames_sync(video_path, tmp / "keys", times))
            logging.info(f"[{label}] distinct 프레임 {len(full)}장 → OCR")
            async with workers._ocr_lock:   # PaddleOCR 인스턴스는 스레드 안전하지 않다
                keyframes = await loop.run_in_executor(
                    None, lambda: ocr_keyframes_sync(full, workers.get_ocr()))

            # 6.타임라인 병합 + 영문 앵커
            timeline = build_timeline(segments, keyframes, long_form)
            anchors, high_anchors = collect_english_anchors(keyframes)
            logging.info(f"[{label}] 6.앵커 {len(anchors)}개 (고신뢰 {len(high_anchors)}개)")

            image_blobs = await loop.run_in_executor(
                None, lambda: prepare_gemini_images(keyframes, GEMINI_IMAGE_BUDGET))

            _job_set(label, entry, 7, "Gemini 합성")
            logging.info(f"[{label}] 7.Gemini 합성 (이미지 {len(image_blobs)}장)")
            await workers.throttle_gemini()
            async with workers.gemini_sem:
                md, usage = await loop.run_in_executor(
                    None, lambda: call_gemini_sync(
                        workers.gemini_client, system_prompt, entry, metadata,
                        timeline, image_blobs, anchors, convert_date_iso))
            if usage:
                await tracker.add(usage)

            # 7b.검증 (규칙 + 커버리지) → 실패 시 1회 재생성
            errs = validate_md(md)
            missing = coverage_missing(md, high_anchors)
            cover_bad = high_anchors and len(missing) > max(5, len(high_anchors) * COVERAGE_RETRY_RATIO)
            if errs or cover_bad:
                note = ""
                if errs:
                    note += f"Rule violations to fix: {'; '.join(errs)}. "
                if cover_bad:
                    note += ("These on-screen English commands / file names / config keys are MISSING "
                             f"from your output and must be integrated verbatim: {', '.join(missing[:25])}.")
                logging.warning(f"[{label}] 검증 보정 재생성: {note[:200]}")
                await asyncio.sleep(5)
                await workers.throttle_gemini()
                async with workers.gemini_sem:
                    md, usage = await loop.run_in_executor(
                        None, lambda: call_gemini_sync(
                            workers.gemini_client, system_prompt, entry, metadata,
                            timeline, image_blobs, anchors, convert_date_iso, retry_note=note))
                if usage:
                    await tracker.add(usage)
                errs = validate_md(md)
                if errs:
                    # 검증 실패해도 삭제하지 않는다. 사람이 검토할 수 있게 보관한다.
                    FAILED_MD_DIR.mkdir(parents=True, exist_ok=True)
                    (FAILED_MD_DIR / filename).write_text(md, encoding="utf-8")
                    logging.warning(f"[{label}] 검증 실패 → {FAILED_MD_DIR.name} 보관: {errs}")
                    log_failure(url, filename, f"재시도 후에도 검증 실패: {errs}")
                    _job_clear(label)
                    return False

            # 8.저장 (result_check\convert_O 평면)
            INBOX_OUT_DIR.mkdir(parents=True, exist_ok=True)
            target_path.write_text(md, encoding="utf-8")
            manifest[url] = {
                "filename": filename,
                "path": str(target_path),
                "krw": round(tracker.krw, 1),
                "ts": datetime.now().isoformat(timespec="seconds"),
            }
            _job_clear(label)
            logging.info(f"[{label}] 저장 완료 → {target_path}")
            logging.info(f"[{label}] {tracker.summary()}")
            return True

        except BudgetExceeded:
            _job_clear(label)
            raise
        except Exception as exc:
            _job_clear(label)
            logging.error(f"[{label}] 예외: {exc}", exc_info=True)
            log_failure(url, filename, str(exc))
            return False


# ──────────────────────── 작업 목록 구성 ──────────────────────────────

def build_work_list(target_entries: list, all_entries: list, convert_date: str):
    """동일 URL은 1회만 변환한다(가이드 섹션 10). 파일을 추가로 만들지 않고 xlsx 행만 전부 O로 갱신한다.
    반환: [(entry, filename, [같은 URL의 모든 xlsx row]), ...]"""
    rows_by_url = {}
    for e in all_entries:
        rows_by_url.setdefault(e["url"], []).append(e["row"])

    work = []
    seen = set()
    for e in target_entries:
        if e["url"] in seen:
            continue
        seen.add(e["url"])
        filename = make_filename(e["channel"], e["video_title"], convert_date)
        work.append((e, filename, rows_by_url.get(e["url"], [e["row"]])))
    return work


# ───────────────────────────── main ──────────────────────────────────

def parse_args():
    ap = argparse.ArgumentParser(description="Tech 유튜브 영상 → .md 변환 (Tech 트랙 전용)")
    ap.add_argument("--dry-run", action="store_true", help="작업 목록만 출력 (API 호출 없음)")
    ap.add_argument("--limit", type=int, help="앞 N개만")
    ap.add_argument("--url", help="특정 URL만")
    ap.add_argument("--urls-file", dest="urls_file", help="처리할 URL 목록 파일 (한 줄에 URL 하나)")
    ap.add_argument("--turbo", action="store_true", help="Whisper large-v3-turbo (속도 우선)")
    ap.add_argument("--max-frames", type=int, help=f"영상당 최대 프레임 (기본 {MAX_FRAMES})")
    ap.add_argument("--model", help=f"Gemini 모델 (기본 {MODEL_NAME})")
    ap.add_argument("--budget-krw", type=float, help="비용 한도 원화 (기본: 무제한)")
    return ap.parse_args()


async def main_async():
    global MODEL_NAME, MAX_FRAMES, BUDGET_KRW
    args = parse_args()
    if args.model:
        MODEL_NAME = args.model
    if args.max_frames:
        MAX_FRAMES = args.max_frames
    if args.budget_krw:
        BUDGET_KRW = args.budget_krw

    setup_logging()
    _active.clear()
    _save_active()
    try:
        from dotenv import load_dotenv
        load_dotenv(BASE_DIR / ".env", override=True)
    except Exception:
        pass

    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key and not args.dry_run:
        logging.error("GEMINI_API_KEY 없음 (.env 확인)")
        sys.exit(1)

    convert_date = datetime.now().strftime("%y%m%d")        # 파일명용 YYMMDD
    convert_date_iso = datetime.now().strftime("%Y-%m-%d")  # Metadata용

    logging.info(f"BASE_DIR = {BASE_DIR}")
    logging.info(f"xlsx 파싱... ({XLSX_PATH})")
    all_entries = parse_xlsx()
    logging.info(f"xlsx 유튜브 행 전체 {len(all_entries)}개")

    # 처리 대상은 X인 행뿐이다. 공란은 X로 간주하지 않는다 (전체 재변환 사고 방지).
    entries = [e for e in all_entries if e["status"] == "X"]
    blank = len([e for e in all_entries if e["status"] == ""])
    logging.info(f"변환 대상(X) {len(entries)}개 | 공란 {blank}개(대상 아님) "
                 f"| O {len([e for e in all_entries if e['status'] == 'O'])}개 "
                 f"| F {len([e for e in all_entries if e['status'] == 'F'])}개")
    if blank and not entries:
        logging.warning("변환완료여부가 전부 공란입니다. 처리할 행을 X로 채워야 스크립트가 인식합니다.")

    if args.url:
        entries = [e for e in entries if e["url"] == args.url]
    if args.urls_file:
        allowed = {u.strip() for u in Path(args.urls_file).read_text(encoding="utf-8").splitlines() if u.strip()}
        entries = [e for e in entries if e["url"] in allowed]

    work = build_work_list(entries, all_entries, convert_date)
    logging.info(f"변환 대상 고유 URL {len(work)}개")
    if args.limit:
        work = work[:args.limit]

    manifest = load_manifest()

    # ── dry-run: API 호출 없음 ──
    if args.dry_run:
        print(f"\n=== DRY RUN ({len(work)}개) ===")
        print(f"xlsx      : {XLSX_PATH}")
        print(f"성공 저장   : {INBOX_OUT_DIR}")
        print(f"실패 보관   : {FAILED_MD_DIR}")
        print(f"모델       : {MODEL_NAME} | 화질 1080p | 최대 프레임 {MAX_FRAMES} | OCR lang={OCR_LANG}")
        if EXEMPLARY_MD and EXEMPLARY_MD.exists():
            print(f"exemplary : {EXEMPLARY_MD}")
        else:
            print("exemplary : 미지정 (규칙만으로 생성. 시범 변환 후 확정 예정)")
        print()
        done = 0
        for entry, fn, rows in work:
            existing = find_existing_md(entry["channel"], entry["video_title"])
            skipped = existing is not None or entry["url"] in manifest
            done += skipped
            mark = "[완료]" if skipped else "[대기]"
            print(f"{mark} [row {entry['row']}] {fn}")
            print(f"   URL : {entry['url']}")
            print(f"   목적 : {entry['purpose']}")
            if len(rows) > 1:
                print(f"   중복 URL → xlsx row {rows} 전부 O로 갱신")
            if existing:
                print(f"   기존 파일: {existing.name}")
        print(f"\n대기 {len(work)-done} | 완료 {done}")

        # 공란 행은 처리 대상이 아니다. 어느 행을 X로 채워야 하는지 보여준다.
        blanks = [e for e in all_entries if e["status"] == ""]
        if blanks:
            print(f"\n--- 공란 {len(blanks)}행 (변환완료여부를 X로 채워야 대상이 됨) ---")
            for e in blanks:
                print(f"  row {e['row']:>3} | {e['channel']} | {e['purpose']} | {e['video_title'][:60]}")
        return

    # ── 실제 실행 ──
    workers = PipelineWorkers()
    if args.turbo:
        workers.whisper_name = "large-v3-turbo"
    workers.init_gemini(api_key)
    system_prompt = build_system_prompt()
    tracker = CostTracker(BUDGET_KRW)

    pipeline_sem = asyncio.Semaphore(PIPELINE_CONC)
    ok = fail = skip = 0
    stop_flag = {"stop": False}

    async def run_one(entry, fn, rows):
        nonlocal ok, fail, skip
        if stop_flag["stop"]:
            return
        url = entry["url"]
        async with pipeline_sem:
            if stop_flag["stop"]:
                return
            # 이미 만든 영상이면 다시 만들지 않는다. xlsx 행만 O로 맞춘다.
            existing = find_existing_md(entry["channel"], entry["video_title"])
            if existing or url in manifest:
                skip += 1
                logging.info(f"건너뜀(이미 존재): {fn}")
                await update_xlsx_status_many(rows, "O")
                return
            try:
                success = await process_video(entry, fn, convert_date_iso,
                                              workers, system_prompt, tracker, manifest)
            except BudgetExceeded as b:
                stop_flag["stop"] = True
                logging.warning(f"비용 한도 도달 → 정지: {b}")
                return
            if success:
                ok += 1
                save_manifest(manifest)   # 성공할 때마다 체크포인트(재개 대비)
                await update_xlsx_status_many(rows, "O")
            else:
                fail += 1
                await update_xlsx_status(entry["row"], "F")

    await asyncio.gather(*[run_one(e, f, r) for e, f, r in work])
    save_manifest(manifest)
    logging.info(f"완료. OK={ok} SKIP={skip} FAIL={fail}")
    logging.info(tracker.summary())
    if stop_flag["stop"]:
        logging.warning("비용 한도로 일부 미처리. 재실행하면 manifest 기준으로 이어서 진행됨.")


def main():
    asyncio.run(main_async())


if __name__ == "__main__":
    main()
